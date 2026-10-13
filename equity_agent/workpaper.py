from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .financials import FinancialRecord


def create_workpaper_excel(
    company_name: str,
    records: list[FinancialRecord],
    result: dict[str, Any],
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["股权投资估值预审工作底稿（初稿）"])
    summary.append(["企业", company_name])
    summary.append(["预审状态", result.get("status", "待复核")])
    summary.append(["重要边界", result.get("scope_note", "仅供估值人员预审，不构成估值结论或投资建议。")])
    summary.append([])
    summary.append(["方法", "低值（亿元）", "基准/高值（亿元）", "说明"])
    valuations = result.get("valuations", {})
    dcf = valuations.get("dcf", {})
    comps = valuations.get("multiples", {})
    summary.append(["DCF 股权价值", _equity(dcf.get("bear")), _equity(dcf.get("bull")), "情景区间；基准情景见 Valuation 页"])
    if comps:
        summary.append(["EV/EBITDA 股权价值", comps.get("low", {}).get("equity_value"), comps.get("high", {}).get("equity_value"), "倍数由使用者录入，不含市场数据验证"])
    summary.append(["提醒", "区间是方法结果的并列展示，不是统计置信区间或正式估值结论。"])

    assumptions = workbook.create_sheet("Assumptions")
    assumptions.append(["假设项", "取值", "口径/说明"])
    assumption_obj = result.get("assumptions")
    assumption_values = vars(assumption_obj) if assumption_obj else {}
    labels = {
        "revenue_growth": ("预测收入年增长率", "%"),
        "fcf_margin": ("预测自由现金流率", "%"),
        "wacc": ("WACC", "%"),
        "terminal_growth": ("永续增长率", "%"),
        "forecast_years": ("预测年数", "年"),
        "multiple_low": ("EV/EBITDA 倍数下限", "x"),
        "multiple_high": ("EV/EBITDA 倍数上限", "x"),
    }
    for key, (label, unit) in labels.items():
        value = assumption_values.get(key)
        if value is not None:
            display = value * 100 if unit == "%" else value
            assumptions.append([label, display, unit + "；由使用者输入，需复核来源"])
    assumptions.append(["净债务", valuations.get("net_debt_used"), "最新年度结构化财务表；需与原始报表核对"])
    assumptions.append(["FCF 历史观测率", _percent(valuations.get("observed_fcf_margin")), "最新年度（经营现金流 - 资本开支）/收入"])
    assumptions.append(["FCF 率本次采用值", _percent(valuations.get("fcf_margin_used")), "用于 DCF 的预测假设"])

    financials = workbook.create_sheet("Financials_Source")
    financials.append(["年度", "营业收入（亿元）", "EBITDA（亿元）", "净利润（亿元）", "经营现金流（亿元）", "资本开支（亿元）", "净债务（亿元）", "来源", "字段单元格/行"])
    for record in records:
        fields = ["revenue", "ebitda", "net_income", "operating_cash_flow", "capex", "net_debt"]
        financials.append([
            record.year,
            *[record.metrics.get(field) for field in fields],
            record.source,
            "; ".join(f"{field}={record.provenance.get(field, '未提供')}" for field in fields),
        ])

    valuation_sheet = workbook.create_sheet("Valuation")
    valuation_sheet.append(["方法/情景", "企业价值（亿元）", "股权价值（亿元）", "补充说明"])
    for scenario, row in dcf.items():
        valuation_sheet.append([f"DCF / {scenario}", row.get("enterprise_value"), row.get("equity_value"), "自由现金流折现；固定公式计算"])
    for name, row in comps.items():
        valuation_sheet.append([f"EV/EBITDA / {name}", row.get("enterprise_value"), row.get("equity_value"), f"输入倍数 {row.get('multiple')}x"])
    if not valuations:
        valuation_sheet.append(["估值未执行", None, None, "存在高优先级缺失项或无效假设；先处理 Review_Queue。"])
    valuation_sheet.append([])
    valuation_sheet.append(["DCF 方法", "先以最新年度收入为基数逐年外推收入和 FCF，再折现显性期 FCF 与永续终值；股权价值 = 企业价值 - 净债务。"])
    valuation_sheet.append(["可比倍数方法", "企业价值 = 最新年度 EBITDA × 使用者输入的 EV/EBITDA 倍数；股权价值 = 企业价值 - 净债务。倍数不是市场数据。"])
    valuation_sheet.append(["情景边界", "Bear/Base/Bull 对增长率、FCF 率和 WACC 做小幅机械扰动，仅用于敏感性观察。"])

    review = workbook.create_sheet("Review_Queue")
    review.append(["级别", "待复核事项"])
    for item in result.get("warnings", []):
        review.append([item.get("severity"), item.get("message")])
    review.append([])
    review.append(["建议动作"])
    for action in result.get("next_actions", []):
        review.append([action])

    documents = workbook.create_sheet("Document_Evidence")
    documents.append(["文件", "处理方式", "页数", "任务 ID", "提示", "识别文本/片段（截取）"])
    for item in result.get("documents", []):
        documents.append([
            item.get("name"), item.get("method"), item.get("pages"), item.get("job_id"),
            item.get("warning"), (item.get("text") or "")[:12000],
        ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else "A1"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="15324B")
        for column in sheet.columns:
            column_letter = get_column_letter(column[0].column)
            max_width = max((len(str(cell.value or "")) for cell in column[:50]), default=10)
            sheet.column_dimensions[column_letter].width = min(max(max_width + 2, 12), 72)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _equity(value: dict[str, Any] | None) -> float | None:
    return value.get("equity_value") if value else None


def _percent(value: float | None) -> float | None:
    return value * 100 if value is not None else None
