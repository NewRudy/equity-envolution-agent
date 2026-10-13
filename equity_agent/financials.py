from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook


METRIC_ALIASES = {
    "revenue": ("revenue", "营业收入", "营业总收入", "主营业务收入"),
    "ebitda": ("ebitda", "息税折旧摊销前利润"),
    "net_income": ("netincome", "净利润", "归母净利润"),
    "operating_cash_flow": (
        "operatingcashflow",
        "经营活动产生的现金流量净额",
        "经营现金流",
    ),
    "capex": (
        "capex",
        "资本开支",
        "购建固定资产无形资产和其他长期资产支付的现金",
    ),
    "net_debt": ("netdebt", "净债务", "净负债", "有息负债减货币资金"),
}
PERIOD_ALIASES = ("year", "年度", "年份", "报告期", "会计年度", "period")


@dataclass
class FinancialRecord:
    year: str
    metrics: dict[str, float | None]
    source: str
    provenance: dict[str, str]


def _header_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[\s（）()【】\[\]：:，,._\-—]+", "", text)
    return re.sub(r"(万亿元|百万元|千元|亿元|万元|元|人民币|单位)$", "", text)


def _unit_scale(value: Any) -> float:
    text = str(value or "")
    if "万亿元" in text:
        return 10000.0
    if "百万元" in text:
        return 0.01
    if "千元" in text:
        return 0.00001
    if "万元" in text:
        return 0.0001
    if "亿元" in text:
        return 1.0
    compact = re.sub(r"[\s（）()]+", "", text)
    if compact.endswith("元"):
        return 0.00000001
    # Template and unlabelled columns are interpreted as CNY 100 million (亿元).
    return 1.0


def _match_header(header: Any) -> tuple[str | None, float]:
    normalized = _header_text(header)
    if normalized in {_header_text(item) for item in PERIOD_ALIASES}:
        return "year", 1.0
    for field, aliases in METRIC_ALIASES.items():
        if normalized in {_header_text(item) for item in aliases}:
            return field, _unit_scale(header)
    return None, 1.0


def _period(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return str(value.year)
    match = re.search(r"(19|20)\d{2}", str(value))
    return match.group(0) if match else None


def _number(value: Any) -> float | None:
    if value is None or str(value).strip() in {"", "-", "—", "--", "不适用", "N/A", "n/a"}:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    text = text.replace("￥", "").replace("¥", "").replace("元", "").replace("亿元", "")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("() ")
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


def parse_financial_file(filename: str, data: bytes) -> list[FinancialRecord]:
    """Parse one tabular financial file and retain cell/row-level provenance."""
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "csv":
        return _parse_csv(filename, data)
    if suffix in {"xlsx", "xlsm"}:
        return _parse_workbook(filename, data)
    raise ValueError("只支持 .xlsx、.xlsm 或 .csv 财务表。")


def _records_from_rows(
    filename: str,
    headers: list[tuple[str | None, float]],
    header_labels: list[Any],
    rows: list[list[Any]],
    locations: list[dict[int, str]],
    sheet_name: str,
) -> list[FinancialRecord]:
    year_col = next((i for i, (field, _) in enumerate(headers) if field == "year"), None)
    revenue_col = next((i for i, (field, _) in enumerate(headers) if field == "revenue"), None)
    if year_col is None or revenue_col is None:
        raise ValueError("未找到年度和营业收入列。请使用模板，或将列名写为：年度、营业收入（亿元）。")

    metric_cols: dict[str, tuple[int, float]] = {}
    for index, (field, scale) in enumerate(headers):
        if field and field not in {"year"} and field not in metric_cols:
            metric_cols[field] = (index, scale)

    output: list[FinancialRecord] = []
    for row_index, row in enumerate(rows):
        if year_col >= len(row):
            continue
        year = _period(row[year_col])
        if not year:
            continue
        metrics: dict[str, float | None] = {}
        provenance: dict[str, str] = {}
        for metric, (column, scale) in metric_cols.items():
            value = row[column] if column < len(row) else None
            number = _number(value)
            metrics[metric] = number * scale if number is not None else None
            if number is not None:
                address = locations[row_index].get(column, f"第 {row_index + 2} 行")
                label = str(header_labels[column] or "") if column < len(header_labels) else ""
                has_unit = any(unit in label for unit in ("万亿元", "百万元", "千元", "亿元", "万元", "元"))
                if not has_unit:
                    address += "（未标单位，按亿元读取）"
                provenance[metric] = address
        output.append(
            FinancialRecord(
                year=year,
                metrics=metrics,
                source=f"{filename} / {sheet_name}",
                provenance=provenance,
            )
        )
    if not output:
        raise ValueError("找到列名，但没有可识别的年度数据行。")
    return sorted(output, key=lambda record: int(record.year))


def _parse_workbook(filename: str, data: bytes) -> list[FinancialRecord]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            matrix = list(sheet.iter_rows(values_only=False))
            for header_index, header_row in enumerate(matrix[:20]):
                values = [cell.value for cell in header_row]
                parsed = [_match_header(value) for value in values]
                fields = {field for field, _ in parsed if field}
                if "year" not in fields or "revenue" not in fields:
                    continue
                headers = parsed
                data_rows = [[cell.value for cell in row] for row in matrix[header_index + 1 :]]
                locations = [
                    {column: sheet.cell(row=header_index + row_index + 2, column=column + 1).coordinate for column in range(len(values))}
                    for row_index in range(len(data_rows))
                ]
                return _records_from_rows(filename, headers, values, data_rows, locations, sheet.title)
    finally:
        workbook.close()
    raise ValueError("没有找到包含“年度”和“营业收入”的工作表。")


def _parse_csv(filename: str, data: bytes) -> list[FinancialRecord]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030")
    rows = list(csv.reader(io.StringIO(text)))
    for header_index, row in enumerate(rows[:20]):
        parsed = [_match_header(value) for value in row]
        fields = {field for field, _ in parsed if field}
        if "year" in fields and "revenue" in fields:
            data_rows = rows[header_index + 1 :]
            locations = [{column: f"第 {header_index + row_index + 2} 行" for column in range(len(row))} for row_index in range(len(data_rows))]
            return _records_from_rows(filename, parsed, row, data_rows, locations, "CSV")
    raise ValueError("CSV 中没有找到包含“年度”和“营业收入”的表头。")


def template_bytes() -> bytes:
    """Return a small workbook template whose monetary unit is explicitly 亿元."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financials"
    headers = [
        "年度", "营业收入（亿元）", "EBITDA（亿元）", "净利润（亿元）",
        "经营现金流（亿元）", "资本开支（亿元）", "净债务（亿元）",
    ]
    sheet.append(headers)
    for row in [
        ["2023A", 8.0, 0.8, 0.4, 0.7, 0.3, 1.5],
        ["2024A", 10.0, 1.2, 0.6, 0.9, 0.35, 1.5],
        ["2025A", 13.0, 1.8, 0.9, 1.3, 0.5, 1.5],
    ]:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="15324B")
    widths = [14, 20, 18, 18, 22, 20, 18]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    notes = workbook.create_sheet("Instructions")
    notes.append(["字段", "说明"])
    notes.append(["单位", "模板中所有金额均为人民币亿元；不确定单位时请勿直接用于估值。"])
    notes.append(["必填", "年度、营业收入；建议提供 EBITDA、净利润、经营现金流、资本开支、净债务。"])
    notes.append(["来源", "上传后系统保留工作表、单元格位置，供估值人员逐项复核。"])
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="15324B")
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 92
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
